from pathlib import Path

from docx import Document

from models.standard_document import StandardDocument
from parsers.base import BaseParser


class DocxParser(BaseParser):
    """
    Parser for DOCX files.
    """

    def parse(self, source: str) -> list[StandardDocument]:
        file_path = Path(source)

        if not file_path.exists():
            raise FileNotFoundError(
                f"File not found: {source}"
            )

        if file_path.suffix.lower() != ".docx":
            raise ValueError("Expected a .docx file")

        document = Document(file_path)

        content = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ).strip()

        return [
            StandardDocument(
                source_type="docx",
                title=file_path.name,
                content=content,
                metadata={
                    "file_name": file_path.name,
                    "extension": file_path.suffix,
                    "size_bytes": file_path.stat().st_size,
                },
            )
        ]

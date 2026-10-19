"""End-to-end parser checks over self-generated fixtures.

These tests used to read repo-root ``datasets/``, which is (was) gitignored,
so the suite was red on any fresh clone while the docstring claimed the
samples were checked in. Fixtures are now built at test time from
dependencies the parsers already require (``pymupdf``, ``python-docx``):
no binary files in git and no external data dependency.
"""

from pathlib import Path

import pymupdf
import pytest
from app.ai.parsers.docx_parser import DocxParser
from app.ai.parsers.pdf_parser import PdfParser
from app.ai.parsers.txt_parser import TxtParser
from app.core.config import settings
from docx import Document

# Written into every fixture so we can assert the parser actually
# round-trips text instead of merely returning something non-empty.
MARKER = "Sourcify parser fixture marker"

# Shape matches a real meeting-transcript export: header block, attendee
# list, speaker turns, and a closing "Meeting ended after" line.
TRANSCRIPT = "\n".join(
    [
        "Meeting Transcript",
        "Transcript Attendees",
        "Alice Sharma, Bob Iyer",
        "Transcript",
        "Alice Sharma: hello there",
        "Bob Iyer: hi, thanks for joining",
        "Meeting ended after 00:15:30",
    ]
)


@pytest.fixture(scope="module")
def txt_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("parsers") / "sample.txt"
    path.write_text(f"{MARKER}\nsecond line\n", encoding="utf-8")

    return path


@pytest.fixture(scope="module")
def pdf_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("parsers") / "sample.pdf"

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), MARKER)
    document.save(str(path))
    document.close()

    return path


@pytest.fixture(scope="module")
def docx_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("parsers") / "sample.docx"

    document = Document()
    document.add_paragraph(MARKER)
    document.save(str(path))

    return path


@pytest.fixture(scope="module")
def transcript_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("parsers") / "transcript.docx"

    document = Document()

    for line in TRANSCRIPT.splitlines():
        document.add_paragraph(line)

    document.save(str(path))

    return path


@pytest.mark.parametrize(
    ("fixture_name", "parser_cls", "source_type"),
    [
        ("txt_path", TxtParser, "txt"),
        ("pdf_path", PdfParser, "pdf"),
        ("docx_path", DocxParser, "docx"),
    ],
)
def test_parser_extracts_content(
    fixture_name: str,
    parser_cls,
    source_type: str,
    request,
):
    path = request.getfixturevalue(fixture_name)

    documents = parser_cls().parse(str(path))

    assert len(documents) == 1

    document = documents[0]

    assert document.source_type == source_type
    assert document.title == path.name
    assert MARKER in document.content
    assert document.metadata["file_name"] == path.name
    assert document.metadata["extension"] == f".{source_type}"
    assert document.metadata["size_bytes"] > 0


def test_pdf_metadata_includes_page_count(pdf_path: Path):
    document = PdfParser().parse(str(pdf_path))[0]

    assert document.metadata["page_count"] == 1


def test_docx_transcript_metadata(transcript_path: Path):
    document = DocxParser().parse(str(transcript_path))[0]

    assert document.metadata["document_type"] == "transcript"
    assert document.metadata["attendees"] == ["Alice Sharma", "Bob Iyer"]
    assert document.metadata["duration"] == "00:15:30"

    # Speaker turns are normalized onto a single line.
    assert "Alice Sharma: hello there" in document.content
    assert "Bob Iyer: hi, thanks for joining" in document.content


@pytest.mark.parametrize("parser_cls", [TxtParser, PdfParser, DocxParser])
def test_missing_file_raises(parser_cls, tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        parser_cls().parse(str(tmp_path / "does-not-exist.txt"))


@pytest.mark.parametrize(
    ("parser_cls", "filename"),
    [
        (TxtParser, "wrong.pdf"),
        (PdfParser, "wrong.txt"),
        (DocxParser, "wrong.txt"),
    ],
)
def test_wrong_extension_raises(parser_cls, filename: str, tmp_path: Path):
    path = tmp_path / filename
    path.write_text("irrelevant", encoding="utf-8")

    with pytest.raises(ValueError, match="Expected a"):
        parser_cls().parse(str(path))


def test_txt_content_over_cap_raises(txt_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "max_extracted_chars", 5)

    with pytest.raises(ValueError, match="content too large"):
        TxtParser().parse(str(txt_path))


def test_pdf_page_cap_raises(pdf_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "max_pdf_pages", 0)

    with pytest.raises(ValueError, match="max 0 pages"):
        PdfParser().parse(str(pdf_path))


def test_docx_empty_raises(tmp_path: Path):
    path = tmp_path / "empty.docx"
    Document().save(str(path))

    with pytest.raises(ValueError, match="No readable text"):
        DocxParser().parse(str(path))

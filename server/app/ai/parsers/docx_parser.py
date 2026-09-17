from pathlib import Path

from docx import Document

from app.ai.models.standard_document import StandardDocument
from app.ai.parsers.base import BaseParser
from app.core.config import settings


class DocxParser(BaseParser):
    """
    Parser for Microsoft Word (.docx) files.
    """

    def parse(
        self,
        source: str,
    ) -> list[StandardDocument]:
        file_path = Path(source)

        if not file_path.exists():
            raise FileNotFoundError(
                f"File not found: {source}"
            )

        if file_path.suffix.lower() != ".docx":
            raise ValueError(
                "Expected a .docx file"
            )

        if file_path.stat().st_size > settings.max_upload_bytes:
            raise ValueError(
                f"File too large: {file_path.name} "
                f"(max {settings.max_upload_bytes // (1024 * 1024)}MB)"
            )

        document = Document(file_path)

        paragraphs = []
        total_chars = 0

        for paragraph in document.paragraphs:
            if not paragraph.text.strip():
                continue

            total_chars += len(paragraph.text)

            if total_chars > settings.max_extracted_chars:
                raise ValueError(
                    f"File content too large: {file_path.name} "
                    f"(max {settings.max_extracted_chars} characters)"
                )

            paragraphs.append(paragraph.text)

        content = "\n".join(paragraphs)

        return [
            StandardDocument(
                source_type="docx",
                title=file_path.name,
                content=content,
                metadata={
                    "file_name": file_path.name,
                    "extension": file_path.suffix.lower(),
                    "size_bytes": file_path.stat().st_size,
                },
            )
        ]
